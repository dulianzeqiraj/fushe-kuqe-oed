"""prepare_data.py
Extract the analysis tables from the primary field workbook and from the USGS
MRVA file into flat CSVs that the MATLAB pipeline reads.

Inputs (read-only, never modified):
  Fushe_Kuqe_All_Data.xlsx      the Desktop copy, which carries three sheets
                                transcribed from Cenameri & Beqiraj (2016)
                                in addition to the well and boundary tables
  Pugh2023_MRVA_44wells_REAL.csv

Outputs (written into ../data):
  wells_all.csv        180 wells, every column of sheet 2
  wells_K.csv          the 43 wells carrying a pumping-test K
  wells_NO3.csv        the 31 wells carrying a nitrate measurement
  boundary.csv         3411 aquifer boundary vertices
  swi_chemistry.csv    4 wells, major ions, Cenameri & Beqiraj (2016) Table 1
  swi_cl_trend.csv     chloride 1984/1999/2001 at wells 341 and 503, Table 2
  swi_facts.csv        flow direction, heads, thicknesses, wedge geometry
  mrva_44.csv          the USGS Mississippi River Valley alluvial set

No value is altered, imputed or rounded here: rows are copied or dropped only.
"""
import csv
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, '..', 'data'))

XLSX = r'C:/Users/d_zeq/OneDrive/Desktop/Fushe_Kuqe_All_Data.xlsx'
MRVA = (r'C:/Users/d_zeq/OneDrive/Desktop/PROJEKTE 2 ARTIKUJ/'
        r'AKUIFER TE DHENA TE NDRYSHME FUSHE KUQE TE DHENA/'
        r'Pugh2023_MRVA_44wells_REAL.csv')


def sheet_rows(wb, name):
    ws = wb[name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return rows


def write_csv(path, header, rows):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in rows:
            w.writerow(['' if v is None else v for v in r])
    return len(rows)


def main():
    os.makedirs(OUT, exist_ok=True)
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

    rows = sheet_rows(wb, '2_All_Wells')
    header = [str(h) for h in rows[0]]
    body = [r for r in rows[1:] if r[0] not in (None, '')]
    n_all = write_csv(os.path.join(OUT, 'wells_all.csv'), header, body)

    iK = header.index('K_m_per_day')
    iN = header.index('NO3_mg_per_L')
    iIn = header.index('Inside_Aquifer')
    iLit = header.index('Lithology')

    withK = [r for r in body if r[iK] not in (None, '')]
    withN = [r for r in body if r[iN] not in (None, '')]
    write_csv(os.path.join(OUT, 'wells_K.csv'), header, withK)
    write_csv(os.path.join(OUT, 'wells_NO3.csv'), header, withN)

    brows = sheet_rows(wb, '3_Boundary_Coords')
    bhead = [str(h) for h in brows[0]]
    bbody = [r for r in brows[1:] if r[0] not in (None, '')]
    n_b = write_csv(os.path.join(OUT, 'boundary.csv'), bhead, bbody)

    # --- the three seawater-intrusion sheets, copied verbatim ---------------
    # Each has two banner lines and a blank line before the real header, so the
    # header row is located rather than assumed.
    swi = {'5_Cenameri2016_Chemistry': 'swi_chemistry.csv',
           '6_Cenameri2016_Cl_Trend': 'swi_cl_trend.csv',
           '7_Cenameri2016_SWI_Facts': 'swi_facts.csv'}
    swi_written = {}
    for sheet, fname in swi.items():
        if sheet not in wb.sheetnames:
            continue
        rs = sheet_rows(wb, sheet)
        hdr = None
        for k, r in enumerate(rs):
            cells = [c for c in r if c not in (None, '')]
            if len(cells) >= 2 and k > 0:
                hdr = k
                break
        body2 = [r for r in rs[hdr + 1:]
                 if any(c not in (None, '') for c in r)
                 and not str(r[0]).startswith(('Note', '- ', 'Notes'))]
        swi_written[fname] = write_csv(os.path.join(OUT, fname),
                                       [str(c) for c in rs[hdr]], body2)

    with open(MRVA, encoding='utf-8') as f:
        mr = list(csv.reader(f))
    write_csv(os.path.join(OUT, 'mrva_44.csv'), mr[0], mr[1:])

    lit = {}
    for r in withK:
        lit[r[iLit]] = lit.get(r[iLit], 0) + 1

    manifest = {
        'source_workbook': XLSX,
        'source_mrva': MRVA,
        'n_wells_total': n_all,
        'n_wells_with_K': len(withK),
        'n_wells_with_NO3': len(withN),
        'n_wells_inside': sum(1 for r in body if r[iIn] == 'YES'),
        'n_inside_with_K': sum(1 for r in withK if r[iIn] == 'YES'),
        'n_inside_with_NO3': sum(1 for r in withN if r[iIn] == 'YES'),
        'n_boundary_vertices': n_b,
        'lithology_of_K_wells': lit,
        'n_mrva': len(mr) - 1,
        'swi_sheets': swi_written,
    }
    with open(os.path.join(OUT, 'manifest.json'), 'w', encoding='utf-8') as f:
        json.dump(manifest, f, indent=2)
    print(json.dumps(manifest, indent=2))


if __name__ == '__main__':
    sys.exit(main())
